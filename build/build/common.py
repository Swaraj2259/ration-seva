"""Shared helpers for building RationSeva .docx documents."""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIAG = os.path.join(ROOT, "diagrams")
OUT = os.path.join(ROOT, "output")
os.makedirs(OUT, exist_ok=True)

NAVY = RGBColor(0x0B, 0x3D, 0x91)


def new_doc(title, subtitle, version="1.0"):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = s.bottom_margin = Inches(0.8)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("RationSeva – Online Ration Card Services Portal")
    r.bold = True; r.font.size = Pt(20); r.font.color.rgb = NAVY
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title); r.bold = True; r.font.size = Pt(16)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"{subtitle}\nDistrict Supply Office | Enterprise e-PDS Platform | Version {version} | 2026").italic = True
    return d


def h(d, text, level=1):
    hd = d.add_heading(text, level=level)
    for r in hd.runs:
        r.font.color.rgb = NAVY
    return hd


def para(d, text, bold_prefix=None):
    p = d.add_paragraph()
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(4)
    return p


def bullets(d, items):
    for it in items:
        p = d.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            p.add_run(it[0]).bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)
        p.paragraph_format.space_after = Pt(1)


def numbered(d, items):
    for it in items:
        p = d.add_paragraph(style="List Number")
        if isinstance(it, tuple):
            p.add_run(it[0]).bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def table(d, header, rows, widths=None, font=9.5, highlight=None):
    """highlight: function(row)->hex fill or None"""
    t = d.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, txt in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(str(txt)); r.bold = True; r.font.size = Pt(font)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(c, "0B3D91")
    for row in rows:
        cells = t.add_row().cells
        fill = highlight(row) if highlight else None
        for i, txt in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(txt)); r.font.size = Pt(font)
            if fill:
                _shade(cells[i], fill)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def image(d, name, width=6.5, caption=None):
    path = os.path.join(DIAG, name)
    if os.path.exists(path):
        d.add_picture(path, width=Inches(width))
        d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        para(d, f"[diagram missing: {name}]")
    if caption:
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(caption); r.italic = True; r.font.size = Pt(9)


def note(d, text):
    p = d.add_paragraph()
    r = p.add_run("Note: " + text); r.italic = True; r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def save(d, fname):
    path = os.path.join(OUT, fname)
    d.save(path)
    print("saved", path)
